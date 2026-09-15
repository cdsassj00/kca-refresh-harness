"""문서 파싱(pdf/hwpx/hwp/docx/txt/md)과 접수(intake). DESIGN.md 6절의 계약을 구현한다.

- parse_document(path) -> (text_md, meta)
  text_md: 페이지 표시 `<!-- page N -->` 가 붙은 Markdown 텍스트
  meta:    {"pages", "chars", "format", "title_guess", "published_guess", "scanned"}
- intake(report_id, path, title=None, published=None) -> meta
  reports/<id>/00_source/<id>.md, 01_meta.json 저장 + registry.csv upsert
"""
from __future__ import annotations

import importlib.util
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Optional

import config

__all__ = ["ParseError", "parse_document", "intake", "guess_from_filename", "SUPPORTED_FORMATS"]

SUPPORTED_FORMATS = ("pdf", "hwpx", "hwp", "docx", "txt", "md")

# 파일명 규칙: NN_[YYYY.MM]_제목.ext  (예: 01_[2023.04]_제목.pdf, 예비06_[2019.08]_제목.pdf)
_FN = re.compile(r"^(?:예비)?(\d{2,3})_\[(\d{4})\.(\d{2})\]_(.+)$")
_PUB = re.compile(r"^(\d{4})[.\-/](\d{1,2})(?:[.\-/]\d{1,2})?$")


class ParseError(Exception):
    """문서를 읽을 수 없을 때. 메시지는 한국어."""


# ---------------------------------------------------------------- 공통 유틸
def guess_from_filename(name: str) -> tuple[str, Optional[str]]:
    """파일명에서 (title_guess, published_guess 'YYYY-MM'|None)를 추정한다."""
    stem = Path(name).stem
    m = _FN.match(stem)
    if m:
        _, y, mo, title = m.groups()
        return title.replace("_", " ").strip(), f"{y}-{mo}"
    return stem.replace("_", " ").strip(), None


def normalize_published(value: Optional[str]) -> Optional[str]:
    """'2023.04' / '2023-4' / '2023-04-20' -> '2023-04'. 형식이 어긋나면 ValueError."""
    if value is None or str(value).strip() == "":
        return None
    m = _PUB.match(str(value).strip())
    if not m:
        raise ValueError(f"발간연월 형식 오류(YYYY-MM 필요): {value}")
    y, mo = int(m.group(1)), int(m.group(2))
    if not (1 <= mo <= 12):
        raise ValueError(f"발간연월의 월이 잘못됨: {value}")
    return f"{y:04d}-{mo:02d}"


def _page_block(n: int, text: str) -> str:
    return f"\n\n<!-- page {n} -->\n{text}"


def _collapse_blank(lines: list[str]) -> list[str]:
    out: list[str] = []
    for ln in lines:
        if ln.strip() == "" and (not out or out[-1].strip() == ""):
            continue
        out.append(ln)
    while out and out[-1].strip() == "":
        out.pop()
    return out


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _sniff_format(path: Path) -> str:
    """확장자 우선, 내용(매직 바이트)으로 보정. 이름만 .hwp 인 HWPX(zip)·PDF 를 잡아낸다."""
    ext = path.suffix.lower().lstrip(".")
    if ext == "markdown":
        ext = "md"
    if ext not in SUPPORTED_FORMATS:
        raise ParseError(f"지원하지 않는 형식: .{ext or '(없음)'} (pdf, hwpx, hwp, docx, txt, md 만 지원)")
    try:
        head = path.open("rb").read(8)
    except OSError as e:
        raise ParseError(f"파일을 열 수 없음: {path.name} ({e})") from e
    if head.startswith(b"%PDF"):
        return "pdf"
    if ext in ("hwp", "hwpx", "docx") and zipfile.is_zipfile(path):
        try:
            with zipfile.ZipFile(path) as z:
                names = set(z.namelist())
        except zipfile.BadZipFile:
            return ext
        if any(n.startswith("Contents/section") for n in names):
            return "hwpx"
        if "word/document.xml" in names:
            return "docx"
    return ext


# ---------------------------------------------------------------- PDF
def _parse_pdf(path: Path) -> tuple[str, dict]:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover - 배포 환경에는 항상 있음
        raise ParseError("PDF 파싱에는 pypdf 설치가 필요합니다 (pip install pypdf)") from e
    try:
        reader = PdfReader(str(path))
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception as e:
                raise ParseError("암호가 걸린 PDF는 열 수 없습니다") from e
        n_pages = len(reader.pages)
    except ParseError:
        raise
    except Exception as e:
        raise ParseError(f"PDF를 읽을 수 없음: {type(e).__name__}: {e}") from e

    blocks: list[str] = []
    empty = 0
    for i, page in enumerate(reader.pages, 1):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        if not t.strip():
            empty += 1
        blocks.append(_page_block(i, t))
    text = "".join(blocks)
    scanned = n_pages > 0 and (empty / n_pages) > 0.5

    title_guess, published_guess = guess_from_filename(path.name)
    if published_guess is None:
        # 파일명 규칙이 없으면 PDF 메타데이터의 제목을 보조로 쓴다(빈 값·기본값은 무시)
        try:
            md_title = (reader.metadata or {}).get("/Title") if reader.metadata else None
            md_title = str(md_title).strip() if md_title else ""
            if md_title and md_title.lower() not in ("untitled", "제목 없음"):
                title_guess = md_title
        except Exception:
            pass
    meta = {"pages": n_pages, "chars": len(text), "format": "pdf",
            "title_guess": title_guess, "published_guess": published_guess, "scanned": scanned}
    return text, meta


# ---------------------------------------------------------------- HWPX
_SECTION_RE = re.compile(r"^Contents/section(\d+)\.xml$", re.I)


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _hwpx_cell_text(tc) -> str:
    """표 셀 안의 모든 문단 텍스트를 공백으로 이어 한 칸으로 만든다."""
    parts: list[str] = []
    for el in tc.iter():
        if _local(el.tag) == "t":
            parts.append("".join(el.itertext()))
    txt = " ".join(p.strip() for p in parts if p and p.strip())
    return txt.replace("|", "｜").replace("\n", " ")


def _hwpx_table_lines(tbl) -> list[str]:
    rows: list[list[str]] = []
    for tr in tbl.iter():
        if _local(tr.tag) != "tr":
            continue
        cells = [_hwpx_cell_text(tc) for tc in tr if _local(tc.tag) == "tc"]
        if cells:
            rows.append(cells)
    if not rows:
        return []
    width = max(len(r) for r in rows)
    lines = [""]
    for i, r in enumerate(rows):
        r = r + [""] * (width - len(r))
        lines.append("| " + " | ".join(r) + " |")
        if i == 0:
            lines.append("|" + "---|" * width)
    lines.append("")
    return lines


def _hwpx_walk(elem, out: list[str], buf: list[str]) -> None:
    """문단(hp:p) 아래를 순서대로 훑어 텍스트는 buf 에, 표는 out 에 바로 쓴다."""
    for child in elem:
        tag = _local(child.tag)
        if tag == "t":
            buf.append("".join(child.itertext()))
        elif tag == "tbl":
            _hwpx_flush(buf, out)
            out.extend(_hwpx_table_lines(child))
        elif tag == "p":  # 글상자 등 subList 안의 문단
            _hwpx_flush(buf, out)
            _hwpx_walk(child, out, buf)
            _hwpx_flush(buf, out)
        elif tag in ("secPr", "ctrl", "linesegarray", "lineSegArray"):
            continue
        else:
            _hwpx_walk(child, out, buf)


def _hwpx_flush(buf: list[str], out: list[str]) -> None:
    if buf:
        txt = "".join(buf).strip()
        out.append(txt)
        buf.clear()


def _parse_hwpx(path: Path) -> tuple[str, dict]:
    try:
        from defusedxml import ElementTree as ET
    except ImportError:  # pragma: no cover
        raise ParseError("HWPX 파싱에는 defusedxml 설치가 필요합니다 (pip install defusedxml)")
    try:
        z = zipfile.ZipFile(path)
    except zipfile.BadZipFile as e:
        raise ParseError(f"HWPX 파일이 손상되었거나 zip 형식이 아닙니다: {path.name}") from e
    with z:
        sections = sorted(
            ((int(m.group(1)), n) for n in z.namelist() if (m := _SECTION_RE.match(n))),
            key=lambda x: x[0])
        if not sections:
            raise ParseError("HWPX 안에 Contents/section*.xml 이 없습니다")
        page_no = 0
        blocks: list[str] = []
        for _, name in sections:
            try:
                root = ET.fromstring(z.read(name))
            except Exception as e:
                raise ParseError(f"HWPX 본문 XML을 읽을 수 없음({name}): {type(e).__name__}") from e
            page_no += 1
            page_lines: list[str] = []
            for p in root:
                if _local(p.tag) != "p":
                    continue
                if page_lines and str(p.get("pageBreak", "0")) in ("1", "true", "True"):
                    blocks.append(_page_block(page_no, "\n".join(_collapse_blank(page_lines))))
                    page_no += 1
                    page_lines = []
                buf: list[str] = []
                before = len(page_lines)
                _hwpx_walk(p, page_lines, buf)
                _hwpx_flush(buf, page_lines)
                if len(page_lines) == before:
                    page_lines.append("")  # 빈 문단
            blocks.append(_page_block(page_no, "\n".join(_collapse_blank(page_lines))))
    text = "".join(blocks)
    title_guess, published_guess = guess_from_filename(path.name)
    meta = {"pages": page_no, "chars": len(text), "format": "hwpx",
            "title_guess": title_guess, "published_guess": published_guess, "scanned": False}
    return text, meta


# ---------------------------------------------------------------- HWP (pyhwp)
def hwp_available() -> bool:
    return bool(shutil.which("hwp5txt")) or importlib.util.find_spec("hwp5") is not None


def _parse_hwp(path: Path) -> tuple[str, dict]:
    if shutil.which("hwp5txt"):
        cmd = [shutil.which("hwp5txt"), str(path)]
    elif importlib.util.find_spec("hwp5") is not None:
        cmd = [sys.executable, "-m", "hwp5.hwp5txt", str(path)]
    else:
        raise ParseError("HWP는 pyhwp 설치 또는 PDF 변환 필요 (pip install pyhwp, 또는 한글에서 PDF/HWPX로 저장)")
    try:
        proc = subprocess.run(cmd, capture_output=True, timeout=300)
    except (OSError, subprocess.TimeoutExpired) as e:
        raise ParseError(f"hwp5txt 실행 실패: {type(e).__name__}: {e}") from e
    if proc.returncode != 0:
        err = proc.stderr.decode("utf-8", errors="replace").strip()[-300:]
        raise ParseError(f"hwp5txt 오류(코드 {proc.returncode}): {err or '원인 불명'}")
    raw = proc.stdout.decode("utf-8", errors="replace")
    # hwp5txt 는 쪽 구분을 주지 않는다. 폼피드가 있으면 그것으로, 없으면 한 쪽으로 취급
    pages = [p for p in raw.split("\f")] if "\f" in raw else [raw]
    text = "".join(_page_block(i, p) for i, p in enumerate(pages, 1))
    title_guess, published_guess = guess_from_filename(path.name)
    meta = {"pages": len(pages), "chars": len(text), "format": "hwp",
            "title_guess": title_guess, "published_guess": published_guess, "scanned": False}
    return text, meta


# ---------------------------------------------------------------- DOCX
def docx_available() -> bool:
    return importlib.util.find_spec("markitdown") is not None or importlib.util.find_spec("docx") is not None


def _parse_docx(path: Path) -> tuple[str, dict]:
    body: Optional[str] = None
    if importlib.util.find_spec("markitdown") is not None:
        try:
            from markitdown import MarkItDown
            body = MarkItDown().convert(str(path)).text_content
        except Exception:
            body = None  # markitdown 실패 시 python-docx 로 재시도
    if body is None and importlib.util.find_spec("docx") is not None:
        try:
            import docx  # python-docx
            d = docx.Document(str(path))
            lines: list[str] = []
            for para in d.paragraphs:
                lines.append(para.text)
            for tbl in d.tables:
                lines.append("")
                for i, row in enumerate(tbl.rows):
                    cells = [c.text.replace("|", "｜").replace("\n", " ").strip() for c in row.cells]
                    lines.append("| " + " | ".join(cells) + " |")
                    if i == 0:
                        lines.append("|" + "---|" * len(cells))
                lines.append("")
            body = "\n".join(_collapse_blank(lines))
        except Exception as e:
            raise ParseError(f"DOCX를 읽을 수 없음: {type(e).__name__}: {e}") from e
    if body is None:
        raise ParseError("DOCX 파싱에는 markitdown 또는 python-docx 설치가 필요합니다")
    text = _page_block(1, body)
    title_guess, published_guess = guess_from_filename(path.name)
    meta = {"pages": 1, "chars": len(text), "format": "docx",
            "title_guess": title_guess, "published_guess": published_guess, "scanned": False}
    return text, meta


# ---------------------------------------------------------------- TXT / MD
def _parse_text(path: Path, fmt: str) -> tuple[str, dict]:
    body = _read_text_file(path).replace("\r\n", "\n")
    text = _page_block(1, body)
    title_guess, published_guess = guess_from_filename(path.name)
    meta = {"pages": 1, "chars": len(text), "format": fmt,
            "title_guess": title_guess, "published_guess": published_guess, "scanned": False}
    return text, meta


# ---------------------------------------------------------------- 진입점
def parse_document(path: Path | str) -> tuple[str, dict]:
    """문서를 페이지 표시가 붙은 Markdown 텍스트와 메타로 바꾼다. 실패하면 ParseError."""
    path = Path(path)
    if not path.is_file():
        raise ParseError(f"파일이 없습니다: {path}")
    fmt = _sniff_format(path)
    if fmt == "pdf":
        return _parse_pdf(path)
    if fmt == "hwpx":
        return _parse_hwpx(path)
    if fmt == "hwp":
        return _parse_hwp(path)
    if fmt == "docx":
        return _parse_docx(path)
    return _parse_text(path, fmt)


def intake(report_id: str, path: Path | str, title: Optional[str] = None,
           published: Optional[str] = None) -> dict:
    """접수: 00_source/<id>.md, 01_meta.json 저장 + registry.csv upsert. 메타를 돌려준다."""
    path = Path(path)
    rdir = config.report_dir(report_id)  # ID 형식 검사 포함
    published_norm = normalize_published(published)
    text, pmeta = parse_document(path)

    src_dir = rdir / "00_source"
    src_dir.mkdir(parents=True, exist_ok=True)
    (src_dir / f"{report_id}.md").write_text(text, encoding="utf-8")

    meta = {
        "report_id": report_id,
        "source_pdf": path.name,
        "pages": pmeta["pages"],
        "chars": pmeta["chars"],
        "published": published_norm or pmeta.get("published_guess"),
        "title": (title or "").strip() or pmeta.get("title_guess") or path.stem,
        "format": pmeta["format"],
        "scanned": bool(pmeta.get("scanned", False)),
    }
    (rdir / "01_meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")

    from scripts.registry import upsert  # CORE_DIR 는 config 가 sys.path 에 넣는다
    upsert(config.REGISTRY_CSV, {
        "id": report_id, "title": meta["title"], "published": meta["published"] or "",
        "source_pdf": meta["source_pdf"], "maturity": "-",
    })
    return meta
